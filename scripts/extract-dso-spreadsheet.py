#!/usr/bin/env python3
"""Extract DSO facts from the mentor spreadsheet into a small TypeScript data file."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "resources" / "dso-documents" / "DSOs 2027.xlsx"
OUTPUT = ROOT / "app" / "dso-spreadsheet-data.ts"

ROW_GROUPS = {
    "andromeda": {"DSOs": [2], "More details": [2]},
    "sombrero": {"DSOs": [3], "More details": [3]},
    "m51": {"DSOs": [4, 5, 6], "More details": [4, 5]},
    "ngc-4536": {"DSOs": [7], "More details": [6]},
    "mcg-07-33-027": {"DSOs": [8], "More details": [7]},
    "ngc-1569": {"DSOs": [9], "More details": [8]},
    "antennae": {"DSOs": [10, 11, 12], "More details": [9, 10]},
    "arp-143": {"DSOs": [13, 14, 15], "More details": [11]},
    "arp-147": {"DSOs": [16, 17, 18], "More details": [12]},
    "cartwheel": {"DSOs": [19], "More details": [13]},
    "m82": {"DSOs": [20, 21, 22], "More details": [14]},
    "gw170817": {"DSOs": [23]},
    "terzan-5": {"DSOs": [24]},
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", "\n").split())


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True)
    result: dict[str, list[dict[str, str]]] = {}
    for key, groups in ROW_GROUPS.items():
        facts: list[dict[str, str]] = []
        for sheet_name, rows in groups.items():
            sheet = workbook[sheet_name]
            headers = [clean(sheet.cell(1, column).value) for column in range(1, sheet.max_column + 1)]
            for row in rows:
                subject = clean(sheet.cell(row, 1).value).lstrip("- ")
                for column, label in enumerate(headers[1:], start=2):
                    value = clean(sheet.cell(row, column).value)
                    if label and value:
                        facts.append({"subject": subject, "label": label, "value": value, "sheet": sheet_name})
        result[key] = facts

    payload = json.dumps(result, ensure_ascii=False, indent=2)
    OUTPUT.write_text(
        "export type SpreadsheetFact = { subject: string; label: string; value: string; sheet: string };\n\n"
        f"export const spreadsheetDsoFacts: Record<string, SpreadsheetFact[]> = {payload};\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUTPUT} with {sum(map(len, result.values()))} facts")


if __name__ == "__main__":
    main()
